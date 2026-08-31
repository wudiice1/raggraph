"""DOCX 文本提取（python-docx）：段落（标题层级识别）+ 表格结构化提取。"""
from pathlib import Path
from docx import Document as DocxDocument


def extract_docx(path: Path) -> str:
    """提取 Word 文档内容，识别标题层级并格式化表格。"""
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        raise ValueError(f"无法读取 Word 文档，文件可能已损坏或格式不兼容: {exc}") from exc

    parts: list[str] = []

    # 1. 提取所有段落并识别标题层级
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        style_name = (p.style.name if p.style and p.style.name else "").lower()
        if "heading 1" in style_name or "标题 1" in style_name:
            parts.append(f"# {text}")
        elif "heading 2" in style_name or "标题 2" in style_name:
            parts.append(f"## {text}")
        elif "heading 3" in style_name or "标题 3" in style_name:
            parts.append(f"### {text}")
        elif "heading 4" in style_name or "标题 4" in style_name:
            parts.append(f"#### {text}")
        else:
            parts.append(text)

    # 2. 提取表格并保持行列结构
    for table in doc.tables:
        table_rows: list[str] = []
        for row in table.rows:
            cell_texts = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            # 过滤全空行
            if any(cell_texts):
                table_rows.append(" | ".join(cell_texts))
        if table_rows:
            parts.append("\n".join(table_rows))

    return "\n\n".join(parts)
